# -*- coding: utf-8 -*-
"""解析 OBIMD 的部件映射表：Sub-character ↔ Glyph Code Point / Main-character"""
import io, json, os, sys, collections
import openpyxl
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

HERE = os.path.dirname(os.path.abspath(__file__))
OB = os.path.join(HERE, "data", "obimd")

for fn in ["Sub-character to Glyph Code Point Mapping.xlsx",
           "Sub-character to Main-character Mapping.xlsx"]:
    p = os.path.join(OB, fn)
    print("=" * 78)
    print(fn)
    print("=" * 78)
    wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    print("  sheets:", wb.sheetnames)
    for ws in wb.worksheets:
        print(f"  sheet '{ws.title}'  {ws.max_row} 行 × {ws.max_column} 列")
        rows = ws.iter_rows(values_only=True)
        for i, r in enumerate(rows):
            print("   ", r)
            if i >= 8:
                break
    wb.close()
    print()

print("=" * 78)
print("README 摘要")
print("=" * 78)
print(open(os.path.join(OB, "README.md"), encoding="utf-8").read()[:2000])
