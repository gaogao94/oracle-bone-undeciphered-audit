# -*- coding: utf-8 -*-
"""
字形家族分析：把 OBIMD 的
  · Sub-character → 字形码（41,687 行，= 每个字形类有多少个具体字形）
  · Sub-character → Main-character（2,748 行，= 异体归并关系）
  · Main-character → 现代字（transcription）
三层连接起来，为每个候选补上「字形」这一类证据。
并核查 README 所称 "five Shang Dynasty phases" 是否可从数据推断。
"""
import io, json, os, sys, collections
import openpyxl
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

HERE = os.path.dirname(os.path.abspath(__file__))
D = os.path.join(HERE, "data")
OB = os.path.join(D, "obimd")
RES = os.path.join(HERE, "result")

# ---- 读表
wb = openpyxl.load_workbook(os.path.join(OB, "Sub-character to Glyph Code Point Mapping.xlsx"),
                            read_only=True, data_only=True)
glyphs = collections.defaultdict(set)
for i, (sub, gcp) in enumerate(wb.active.iter_rows(values_only=True)):
    if i == 0 or not sub:
        continue
    glyphs[sub].add(gcp)
wb.close()

wb = openpyxl.load_workbook(os.path.join(OB, "Sub-character to Main-character Mapping.xlsx"),
                            read_only=True, data_only=True)
sub2main = {}
for i, (sub, main) in enumerate(wb.active.iter_rows(values_only=True)):
    if i == 0 or not sub:
        continue
    sub2main[sub] = main
wb.close()

mc = json.load(open(os.path.join(OB, "main_character.json"), encoding="utf-8"))
data = json.load(open(os.path.join(OB, "data.json"), encoding="utf-8"))

print("=" * 80)
print("字形层统计")
print("=" * 80)
print(f"  Sub-character（子类）      {len(glyphs):,}")
print(f"  字形码总数（去重）          {len({g for s in glyphs.values() for g in s}):,}")
print(f"  sub→main 映射              {len(sub2main):,}")
sizes = collections.Counter(len(v) for v in glyphs.values())
print(f"  每个子类的字形数分布（前 12）: {dict(sorted(sizes.items())[:12])}")
allg = [len(v) for v in glyphs.values()]
print(f"  子类字形数：均值 {sum(allg)/len(allg):.1f}，中位 {sorted(allg)[len(allg)//2]}，最大 {max(allg)}")

# ---- 把字形数聚合到 Main-character（Label）
main_glyphs = collections.defaultdict(set)
main_subs = collections.defaultdict(set)
for sub, gs in glyphs.items():
    m = sub2main.get(sub, sub)
    main_glyphs[m] |= gs
    main_subs[m].add(sub)

freq = collections.Counter()
for piece in data:
    for grp in piece.get("RecordUtilSentenceGroupVoList") or []:
        for c in grp.get("RecordUtilOracleCharVoList") or []:
            if c.get("Label"):
                freq[c["Label"]] += 1

print(f"\n  覆盖到的 Main-character  {len(main_glyphs):,}")
print(f"  语料中出现的 Label        {len(freq):,}")

# ---- 断代线索：README 称五个商代时期
print("\n" + "=" * 80)
print("断代线索核查：数据里有没有'期'的信息？")
print("=" * 80)
keys = set()
for p in data[:3000]:
    keys |= set(p.keys())
print(f"  片级字段: {sorted(keys)}")
print("  → 无 period/phase 字段。README 的 'across five Shang Dynasty phases'")
print("    应指数据集覆盖五个时期，但未逐片标注。类组只能从外部补。")

# ---- 为候选补字形证据
cands = json.load(open(os.path.join(RES, "hapax_solvability_v3.json"), encoding="utf-8"))
hyp = json.load(open(os.path.join(RES, "hypothesis_check.json"), encoding="utf-8"))
hypmap = {h["glyph"]: h for h in hyp}

print("\n" + "=" * 80)
print("26 条假设的字形证据补充")
print("=" * 80)
print(f"{'未释字':>4} {'Lab':>12} {'出现':>4} {'异构体':>5} {'字形数':>5} {'假设读法':<18} {'完整串':>5}")
print("-" * 80)
enriched = []
label_by_glyph = {}
for r in cands["rows"]:
    if r["occurrences"] == 1 and r["parallel"]:
        label_by_glyph[r["glyph"]] = r["label"]

for g, lab in label_by_glyph.items():
    ng = len(main_glyphs.get(lab, ()))
    ns = len(main_subs.get(lab, ()))
    h = hypmap.get(g, {})
    print(f"{g:>4} {lab:>12} {freq[lab]:>4} {ns:>5} {ng:>5} {h.get('hypothesis','—'):<18} {h.get('full','—'):>5}")
    enriched.append({"glyph": g, "label": lab, "occurrences": freq[lab],
                     "n_sublabels": ns, "n_glyphs": ng,
                     "hypothesis": h.get("hypothesis"), "full_corpus_hits": h.get("full")})

# ---- 已释字的字形数基准
known_ng = [len(main_glyphs[l]) for l in freq
            if (mc.get(l, {}).get("transcription") or []) and l in main_glyphs]
known_fq = [freq[l] for l in freq
            if (mc.get(l, {}).get("transcription") or []) and l in main_glyphs]
print(f"\n  已释字（有隶定且有字形数据）{len(known_ng)} 个：字形数均值 "
      f"{sum(known_ng)/len(known_ng):.1f}，中位 {sorted(known_ng)[len(known_ng)//2]}")
print(f"  匿名候选字形数均值 "
      f"{sum(e['n_glyphs'] for e in enriched)/len(enriched):.1f}")

json.dump({"enriched": enriched,
           "stats": {"n_subchar": len(glyphs), "n_glyphcodes": len({g for s in glyphs.values() for g in s}),
                     "known_glyph_mean": sum(known_ng)/len(known_ng)}},
          open(os.path.join(RES, "glyph_evidence.json"), "w", encoding="utf-8"),
          ensure_ascii=False, indent=1)
print("\n[写出] " + os.path.join(RES, "glyph_evidence.json"))
